"""Verificación global del catálogo, las fichas, los CSV de salida y el libro."""
import csv
import re
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
CATALOGO = RAIZ / "etl" / "catalogo"
SALIDA = RAIZ / "etl" / "salida"
FICHAS = RAIZ / "docs" / "fuentes"
LIBRO = RAIZ / "Data Dils_v2.xlsx"

# Tarea 15 cancelada por el usuario: siguen solo_diseno, sin esquema, tabla ni ficha.
SIN_TABLA = {"NOT-NP-MENSUAL", "NOT-PEN", "IDEALISTA-ALQUILER", "IDEALISTA-VENTA"}
ERRORES_EXCEL = {"#REF!", "#VALUE!", "#N/A", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}


def _tablas() -> pd.DataFrame:
    return pd.read_csv(CATALOGO / "tablas.csv", dtype=str, keep_default_na=False)


def _con_tabla() -> pd.DataFrame:
    t = _tablas()
    return t[(t["esquema"] != "") & (t["tabla"] != "")]


def test_todas_las_tablas_en_catalogo():
    t = _tablas()
    assert len(t) == 23
    sin = set(t[(t["esquema"] == "") | (t["tabla"] == "")]["dataset_id"])
    assert sin == SIN_TABLA
    assert set(t[t["dataset_id"].isin(SIN_TABLA)]["estado"]) == {"solo_diseno"}
    sql = (RAIZ / "sql" / "create_tables.sql").read_text(encoding="utf-8")
    creadas = set(re.findall(r"CREATE TABLE IF NOT EXISTS `(\w+)`\.`(\w+)`", sql))
    esperadas = set(zip(_con_tabla()["esquema"], _con_tabla()["tabla"]))
    assert esperadas <= creadas, f"sin CREATE TABLE: {sorted(esperadas - creadas)}"


def test_csv_coinciden_con_campos():
    campos = pd.read_csv(CATALOGO / "campos.csv", dtype=str, keep_default_na=False)
    csvs = sorted(SALIDA.glob("*.csv"))  # el de SERPAVI está en .gitignore: solo si existe
    assert csvs, "no hay CSV en etl/salida"
    for ruta in csvs:
        esperado = list(campos[campos["tabla"] == ruta.stem]["campo"])
        assert esperado, f"{ruta.stem} no está en campos.csv"
        with open(ruta, encoding="utf-8", newline="") as f:
            cabecera = next(csv.reader(f, delimiter=";"))
        assert cabecera == esperado, ruta.name


def test_fichas_existen():
    ids = list(_con_tabla()["dataset_id"])
    assert len(ids) == 19
    faltan = [i for i in ids if not (FICHAS / f"{i}.md").is_file()]
    assert not faltan, f"sin ficha: {faltan}"
    assert not [i for i in SIN_TABLA if (FICHAS / f"{i}.md").exists()]


def test_glosario_completo():
    g = pd.read_csv(CATALOGO / "glosario.csv", dtype=str, keep_default_na=False)
    assert list(g.columns) == ["termino", "definicion", "fuente"]
    assert set(g["fuente"]) <= {"MIVAU", "INE", "proyecto"}
    assert (g["definicion"].str.len() > 0).all() and g["termino"].is_unique
    for t in ["VPO", "Calificación definitiva", "Vivienda libre", "ETI", "Valor tasado", "Precio asking",
              "SERPAVI", "Sección censal", "Código ES", "Código CA", "Código PR", "Código MU", "Código SC",
              "Formato de periodo", "Valores de nivel"]:
        assert t in set(g["termino"]), t


def test_libro_sin_errores():
    if not LIBRO.exists():
        pytest.skip("Data Dils_v2.xlsx no existe (se genera con construir_libro)")
    wb = load_workbook(LIBRO, read_only=True)
    try:
        malas = []
        for ws in wb.worksheets:  # incluye hojas ocultas
            for fila in ws.iter_rows(values_only=True):
                for v in fila:
                    if isinstance(v, str) and v.startswith("#") and v.strip() in ERRORES_EXCEL:
                        malas.append((ws.title, v))
        assert not malas, malas[:10]
    finally:
        wb.close()
