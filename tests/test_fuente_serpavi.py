"""MIVAU-SERPAVI contra la muestra versionada tests/fixtures/serpavi_muestra.xlsx.

La muestra (etl/herramientas/serpavi_muestra.py) copia filas concretas del archivo completo (no versionado) con
su cabecera real y todas las columnas. Las coordenadas de los comentarios son las del ARCHIVO COMPLETO
(2026-03_09_bd_SERPAVI_2011-2024 - DEFINITIVO WEB_v2.xlsx); la ficha docs/fuentes/MIVAU-SERPAVI.md da la fila
de la muestra que corresponde a cada una.
"""
import warnings
from datetime import date
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from etl.fuentes import REGISTRO, mivau_serpavi as serpavi
from etl.lib import secciones_serpavi_generar as generar_secciones
from etl.lib.fuente_comun import LIBRO, validar_o_fallar
from etl.lib.territorios import TerritorioDesconocido, cargar_maestro, territorio_por_codigo

MUESTRA = Path(__file__).parent / "fixtures" / "serpavi_muestra.xlsx"
FECHA = date(2026, 10, 29)


@pytest.fixture(scope="module")
def resultado():
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        tabla = serpavi.transformar(MUESTRA, FECHA)
    return tabla, [str(a.message) for a in avisos]


@pytest.fixture(scope="module")
def tabla(resultado):
    return resultado[0].set_index("id_registro", drop=False)


def _v(tabla, id_registro, medida):
    v = tabla.loc[id_registro, medida]
    return None if pd.isna(v) else v


def test_registrada():
    assert REGISTRO["MIVAU-SERPAVI"] is serpavi
    assert serpavi.TABLA == "mivau_serpavi_alquiler" and serpavi.DIMENSIONES == ["tipologia"]
    assert serpavi.MEDIDAS == ["renta_mediana_eur_m2_mes", "renta_p25_eur_m2_mes", "renta_p75_eur_m2_mes",
                               "renta_mediana_eur_mes", "renta_p25_eur_mes", "renta_p75_eur_mes",
                               "superficie_mediana_m2", "superficie_p25_m2", "superficie_p75_m2",
                               "num_inmuebles_alquiler"]
    assert serpavi.CUADRE is False  # medianas y percentiles no se suman (los recuentos se comprueban aparte)


def test_valores_conocidos_ccaa(tabla):
    # Hoja "CCAA", fila 19 (Ceuta, código '18')
    assert _v(tabla, "CA18_COLECTIVA_2011", "num_inmuebles_alquiler") == 978       # C19
    assert _v(tabla, "CA18_COLECTIVA_2011", "renta_mediana_eur_m2_mes") == 7.69    # E19 = 7.6923076923…
    assert _v(tabla, "CA18_COLECTIVA_2024", "renta_mediana_eur_m2_mes") == 9.45    # JE19 = 9.4488440860…
    assert _v(tabla, "CA18_UNIFAMILIAR_2024", "num_inmuebles_alquiler") == 382     # JD19


def test_valores_conocidos_seccion(tabla):
    # Hoja "Secciones censales", fila 36195 (sección 5100101001 de Ceuta): primer y último año
    sc = "SC5100101001_COLECTIVA_"
    assert _v(tabla, sc + "2011", "num_inmuebles_alquiler") == 44                  # F36195
    assert _v(tabla, sc + "2011", "renta_mediana_eur_m2_mes") == 8.70              # H36195 = 8.7024087…
    assert _v(tabla, sc + "2024", "renta_mediana_eur_m2_mes") == 11.86             # JH36195 = 11.8597800…
    assert _v(tabla, sc + "2024", "renta_p25_eur_m2_mes") == 8.89                  # JI36195 = 8.8888888…
    assert _v(tabla, sc + "2024", "renta_p75_eur_m2_mes") == 14.86                 # JJ36195 = 14.8648648…
    assert _v(tabla, sc + "2024", "renta_mediana_eur_mes") == 751.28               # JN36195 = 751.2781808…
    assert _v(tabla, sc + "2024", "superficie_mediana_m2") == 73                   # JT36195
    assert _v(tabla, sc + "2024", "num_inmuebles_alquiler") == 107                 # JF36195
    # Unifamiliar 2024: recuento 0 literal (JG36195) y renta vacía (JK36195) → 0 se queda 0, la renta NULL
    assert _v(tabla, "SC5100101001_UNIFAMILIAR_2024", "num_inmuebles_alquiler") == 0
    assert _v(tabla, "SC5100101001_UNIFAMILIAR_2024", "renta_mediana_eur_m2_mes") is None


def test_valores_conocidos_municipio_y_distrito(tabla):
    assert _v(tabla, "MU52001_UNIFAMILIAR_2024", "renta_mediana_eur_mes") == 500   # Municipios JP8100 (Melilla)
    assert _v(tabla, "MU52001_UNIFAMILIAR_2024", "num_inmuebles_alquiler") == 540  # Municipios JF8100
    assert _v(tabla, "DI5200108_COLECTIVA_2015", "superficie_mediana_m2") == 73    # Distritos CV10512
    assert _v(tabla, "DI5200108_COLECTIVA_2015", "num_inmuebles_alquiler") == 631  # Distritos CH10512


def test_columnas_de_territorio(tabla):
    s = tabla.loc["SC5100101001_COLECTIVA_2024"]
    assert (s["nivel"], s["cod_ccaa"], s["nombre_ccaa"], s["cod_provincia"], s["cod_municipio"],
            s["nombre_municipio"], s["cod_distrito"], s["cod_seccion"]) == (
        "Seccion", "18", "Ceuta", "51", "51001", "Ceuta", "5100101", "5100101001")
    d = tabla.loc["DI5200108_COLECTIVA_2015"]
    assert d["nivel"] == "Distrito" and d["cod_distrito"] == "5200108" and pd.isna(d["cod_seccion"])
    m = tabla.loc["MU52001_UNIFAMILIAR_2024"]
    assert m["nivel"] == "Municipio" and pd.isna(m["cod_distrito"]) and pd.isna(m["cod_seccion"])
    a = tabla.loc["CA18_COLECTIVA_2024"]
    assert (a["periodo"], a["frecuencia"], a["anio"], str(a["fecha"])) == ("2024", "A", 2024, "2024-01-01")
    assert pd.isna(a["trimestre"]) and pd.isna(a["mes"])


def test_forma(resultado):
    t = resultado[0]
    assert t["id_registro"].is_unique
    assert set(t["nivel"]) == {"CCAA", "Provincia", "Municipio", "Distrito", "Seccion"}
    assert set(t["tipologia"]) == {"COLECTIVA", "UNIFAMILIAR"}  # el archivo no trae total de tipologías
    assert sorted(t["anio"].unique()) == list(range(2011, 2025))
    assert t["provisional"].eq(0).all()
    # Filas (territorio, año, tipología) con alguna de las 10 medidas; las vacías no entran
    assert not t[serpavi.MEDIDAS].isna().all(axis=1).any()
    # el oráculo cuenta también el bloque de Valmadrid 2022 colectiva, cuyo único dato es el recuento no entero
    # (pasa a NULL y el bloque queda vacío: no entra)
    assert len(t) == serpavi_filas_con_dato() - 1


def serpavi_filas_con_dato() -> int:
    """Oráculo independiente: cuenta en la muestra los bloques (fila con código, año, tipología) no vacíos."""
    wb = load_workbook(MUESTRA, read_only=True, data_only=True)
    n = 0
    for hoja in ("CCAA", "Provincias", "Municipios", "Distritos", "Secciones censales"):
        filas = wb[hoja].iter_rows(values_only=True)
        cab = next(filas)
        for f in filas:
            if f[0] is None:
                continue  # fila sin código (CCAA, fila 21)
            for anio in range(11, 25):
                for tip in ("VC", "VU"):
                    cols = [j for j, c in enumerate(cab) if str(c).endswith(f"{tip}_{anio}")
                            or str(c) == f"BI_ALVHEPCO_T{tip}_{anio}"]
                    assert len(cols) == 10
                    if any(f[j] not in (None, "") for j in cols):
                        n += 1
    wb.close()
    return n


def test_fila_sin_codigo_no_entra_y_es_la_suma_de_las_ccaa(resultado):
    # CCAA fila 21: sin código ni nombre, solo recuentos 2012-2024 (W21 = 1060264 …). No se inventa la etiqueta
    # "Nacional": no se carga (pendiente de revisión por el usuario). Comprobación: es la suma de las 19 CCAA.
    t = resultado[0]
    assert "ES" not in set(t["cod_territorio"])
    wb = load_workbook(MUESTRA, read_only=True, data_only=True)
    filas = list(wb["CCAA"].iter_rows(values_only=True))
    wb.close()
    cab, sin_codigo = filas[0], [f for f in filas[1:] if f[0] is None]
    assert len(sin_codigo) == 1 and sin_codigo[0][cab.index("BI_ALVHEPCO_TVC_12")] == 1060264  # W21
    ccaa = t[t["nivel"] == "CCAA"]
    for anio in range(2012, 2025):
        for tip, cod in (("COLECTIVA", "VC"), ("UNIFAMILIAR", "VU")):
            suma = ccaa[(ccaa["anio"] == anio) & (ccaa["tipologia"] == tip)]["num_inmuebles_alquiler"].sum()
            assert suma == sin_codigo[0][cab.index(f"BI_ALVHEPCO_T{cod}_{anio % 100}")]


def test_codigo_manda_sobre_columnas_incoherentes(tabla, resultado):
    # Secciones censales fila 8576: CPRO '12', CUMUN '11013', CUSEC '1101208012' → municipio 11012 (Cádiz)
    # (la fila no tiene datos, así que no entra en la tabla; su código sí está en el maestro, con Cádiz)
    assert not (tabla["cod_territorio"] == "SC1101208012").any()
    s = territorio_por_codigo(cargar_maestro(LIBRO), "SC1101208012")
    assert (s["cod_municipio"], s["nombre_municipio"], s["cod_provincia"]) == ("11012", "Cádiz", "11")
    # Distritos fila 6146: CUMUN '31002' (Abárzuza) pero CUDIS '3100101' → Abáigar (31001); sus recuentos
    # coinciden con los del municipio 31001 y su sección 3100101001
    d = tabla[tabla["cod_territorio"] == "DI3100101"]
    assert len(d) and set(d["cod_municipio"]) == {"31001"}
    assert _v(tabla, "DI3100101_UNIFAMILIAR_2023", "num_inmuebles_alquiler") == \
        _v(tabla, "MU31001_UNIFAMILIAR_2023", "num_inmuebles_alquiler") == 1
    avisos = " ".join(resultado[1])
    assert "1101208012" in avisos and "3100101" in avisos


def test_codigo_incoherente_nuevo_falla():
    with pytest.raises(ValueError, match="9999999999"):
        serpavi.comprobar_coherencia("Secciones censales", "9999999999", {"CPRO": "99", "CUMUN": "99998"})


def test_municipio_extinguido_sin_datos_no_entra(tabla):
    # Municipios fila 2152: Cesuras (15026), fusionado en 2013; fila sin ningún dato
    assert not tabla["cod_territorio"].isin(["MU15026"]).any()


def test_recuento_no_entero_valmadrid_es_nulo(tabla, resultado):
    # Municipios HQ8072 = 103.38323643… (recuento no entero; su distrito y su sección dicen 0)
    # su único dato en el bloque era ese recuento: el bloque queda vacío y no entra
    assert "MU50275_COLECTIVA_2022" not in tabla.index
    assert _v(tabla, "MU50275_UNIFAMILIAR_2022", "num_inmuebles_alquiler") == 1  # HR8072
    assert _v(tabla, "DI5027501_COLECTIVA_2022", "num_inmuebles_alquiler") == 0
    assert any("50275" in a and "no entero" in a for a in resultado[1])


def test_recuento_no_entero_nuevo_falla():
    with pytest.raises(ValueError, match="no entero"):
        serpavi.a_recuento(12.5, ("Municipio", "28079", 2020, "COLECTIVA"))
    assert serpavi.a_recuento(12.0, ("Municipio", "28079", 2020, "COLECTIVA")) == 12
    assert serpavi.a_recuento(None, ("Municipio", "28079", 2020, "COLECTIVA")) is None


def test_cabecera_desconocida_falla():
    cab = ["CCAA", "LITCCAA", "BI_ALVHEPCO_TVC_11", "RARO_VC_11"]
    with pytest.raises(ValueError, match="RARO_VC_11"):
        serpavi.columnas_medidas(cab, 2)


def test_validacion_limpia(resultado):
    validar_o_fallar(resultado[0], serpavi.TABLA, cargar_maestro(LIBRO))


def test_cuadre_recuentos_con_excepciones_exactas(resultado):
    # En la muestra, las jerarquías completas (Ceuta, Melilla, Jaurrieta, Abáigar, Valmadrid) cuadran salvo
    # Jaurrieta 2022 colectiva: el distrito dice 3 y municipio y sección 1 (Navarra 2022, recuentos desplazados).
    vistas = serpavi.diferencias_recuentos(resultado[0])
    assert vistas == {("Municipio", "31134", 2022, "COLECTIVA"): -2.0, ("Distrito", "3113401", 2022, "COLECTIVA"): 2.0}
    assert all(serpavi.EXCEPCIONES_RECUENTOS[k] == v for k, v in vistas.items())
    assert len(serpavi.EXCEPCIONES_RECUENTOS) == 35


def test_cuadre_recuentos_diferencia_nueva_falla(resultado):
    t = resultado[0].copy()
    malo = t["id_registro"] == "SC5100101001_COLECTIVA_2020"
    t.loc[malo, "num_inmuebles_alquiler"] = t.loc[malo, "num_inmuebles_alquiler"] + 1
    with pytest.raises(ValueError, match="5100101"):
        serpavi.comprobar_recuentos(t)


def test_generar_secciones_desde_la_muestra(tmp_path):
    df = generar_secciones.generar(MUESTRA, tmp_path / "s.csv")
    sc, di = df[df["nivel"] == "Seccion"], df[df["nivel"] == "Distrito"]
    assert len(sc) == 105 and len(di) == 19  # 14 de Ceuta/Melilla, 3 de filas sueltas y 2 implícitos en secciones
    f = sc[sc["cod_territorio"] == "SC1101208012"].iloc[0]
    assert (f["cod_municipio"], f["nombre_municipio"], f["cod_provincia"], f["nombre_territorio"]) == (
        "11012", "Cádiz", "11", "Cádiz - distrito 08 - sección 012")


def test_generar_secciones_municipio_desconocido_falla(tmp_path):
    wb = Workbook()
    wb.active.title = "Distritos"
    wb["Distritos"].append(["CPRO", "LITPRO", "CUMUN", "LITMUN", "CUDIS"])
    ws = wb.create_sheet("Secciones censales")
    ws.append(["CPRO", "LITPRO", "CUMUN", "LITMUN", "CUSEC"])
    ws.append(["99", "X", "99999", "X", "9999901001"])
    wb.save(tmp_path / "x.xlsx")
    with pytest.raises(ValueError, match="99999"):
        generar_secciones.generar(tmp_path / "x.xlsx", tmp_path / "s.csv")


def test_codigo_desconocido_en_transformar_falla(tmp_path):
    # Una fila con datos cuyo código no está en el maestro hace fallar el proceso (no se salta)
    wb = load_workbook(MUESTRA)
    ws = wb["Secciones censales"]
    ws.cell(ws.max_row, 5).value = "5100199999"
    ws.cell(ws.max_row, 3).value = "51001"
    ws.cell(ws.max_row, 1).value = "51"
    wb.save(tmp_path / "m.xlsx")
    with pytest.raises(TerritorioDesconocido, match="SC5100199999"):
        serpavi.transformar(tmp_path / "m.xlsx", FECHA, cuadre=False)
