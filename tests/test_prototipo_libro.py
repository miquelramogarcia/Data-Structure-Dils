"""Gancho de hojas "(prototipo)" del libro: el libro de producción no cambia si no se piden."""
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from etl.construir_libro import construir_libro
from prototipos.libro import hojas_prototipo

RAIZ = Path(__file__).resolve().parent.parent
ORIGEN = RAIZ / "Data Dils.xlsx"
pytestmark = pytest.mark.skipif(not ORIGEN.exists(), reason="Falta Data Dils.xlsx")


def _catalogo():
    tablas = pd.DataFrame([{"dataset_id": "D0", "esquema": "MIVAU", "tabla": "t0", "hoja": "", "estadistica": "x",
                            "origen": "o", "fuente": "f", "frecuencia": "Mensual", "formato_origen": "XLS",
                            "prioridad": "1", "estado": "con_datos"}])
    campos = pd.DataFrame([{"tabla": "t0", "campo": c, "tipo_mysql": t, "nulo": "NO", "clave": "", "descripcion": "d",
                            "ejemplo": ""} for c, t in (("cod_territorio", "VARCHAR(8)"), ("valor", "INT"))])
    return tablas, campos


def _construir(tmp_path, nombre, **kw):
    pd.DataFrame({"cod_territorio": ["ES"], "valor": [1]}).to_csv(tmp_path / "t0.csv", sep=";", index=False)
    ruta = tmp_path / nombre
    construir_libro(ORIGEN, ruta, dir_salida=tmp_path, catalogo=_catalogo(), **kw)
    return load_workbook(ruta)


def test_sin_hojas_extra_el_libro_no_cambia(tmp_path):
    base = _construir(tmp_path, "a.xlsx")
    extra = _construir(tmp_path, "b.xlsx", hojas_extra=[("X (prototipo)", pd.DataFrame({"a": [1]}))])
    assert base.sheetnames == ["Tablas", "Campos", "Maestro Territorios", "Glosario", "General", "t0"]
    assert extra.sheetnames == base.sheetnames + ["X (prototipo)"]


def test_hoja_extra_repetida_o_larga_falla(tmp_path):
    df = pd.DataFrame({"a": [1]})
    for nombre in ("Tablas", "x" * 32):
        with pytest.raises(ValueError):
            _construir(tmp_path, "c.xlsx", hojas_extra=[(nombre, df)])


def test_hojas_de_los_prototipos_llevan_sufijo():
    hojas = hojas_prototipo()
    assert hojas, "no hay hojas de prototipo"
    nombres = [n for n, _ in hojas]
    assert all(n.endswith("(prototipo)") and len(n) <= 31 for n in nombres) and len(set(nombres)) == len(nombres)
