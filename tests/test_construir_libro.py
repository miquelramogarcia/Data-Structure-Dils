from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from etl import construir_libro as cl
from etl.construir_libro import construir_libro, nombre_hoja

ORIGEN = Path(__file__).parent.parent / "Data Dils.xlsx"
VPO = "mivau_vpo_calificaciones_definitivas"
pytestmark = pytest.mark.skipif(not ORIGEN.exists(), reason="Falta Data Dils.xlsx")


@pytest.fixture(scope="module")
def libro_real(tmp_path_factory):
    ruta = tmp_path_factory.mktemp("libro") / "v2.xlsx"
    construir_libro(ORIGEN, ruta)
    return load_workbook(ruta)


def _catalogo_sintetico(nombres):
    tablas = pd.DataFrame([{"dataset_id": f"D{i}", "esquema": "MIVAU", "tabla": n, "hoja": "",
                            "estadistica": "x", "origen": "o", "fuente": "f", "frecuencia": "Mensual",
                            "formato_origen": "XLS", "prioridad": "1", "estado": "con_datos"}
                           for i, n in enumerate(nombres)])
    campos = pd.DataFrame([{"tabla": n, "campo": c, "tipo_mysql": t, "nulo": "NO", "clave": "",
                            "descripcion": "d", "ejemplo": ""}
                           for n in nombres for c, t in (("cod_territorio", "VARCHAR(8)"), ("valor", "INT"))])
    return tablas, campos


def _csv(dir_, tabla, filas):
    pd.DataFrame({"cod_territorio": ["ES"] * filas, "valor": range(filas)}).to_csv(
        Path(dir_) / f"{tabla}.csv", sep=";", index=False, encoding="utf-8")


def test_nombre_hoja_recorta_y_desempata():
    usados = set()
    a = nombre_hoja("x" * 40, usados)
    usados.add(a)
    b = nombre_hoja("x" * 45, usados)
    assert len(a) == 31 and len(b) == 31 and a != b and b.endswith("~2")


def test_hipervinculo_apunta_a_hoja(libro_real):
    ws = libro_real["Tablas"]
    fila = next(r for r in ws.iter_rows(min_row=2) if r[1].value == VPO)
    hoja = libro_real.sheetnames[5]
    assert fila[1].hyperlink.location == f"'{hoja}'!A1"


def test_hojas_en_orden(libro_real):
    assert libro_real.sheetnames[:5] == ["Tablas", "Campos", "Maestro Territorios", "Glosario", "General"]
    assert len(libro_real.sheetnames) == 6


def test_nombres_hoja_unicos(tmp_path):
    n1, n2 = "a" * 35 + "_uno", "a" * 35 + "_dos"
    _csv(tmp_path, n1, 3)
    _csv(tmp_path, n2, 3)
    salida = tmp_path / "v2.xlsx"
    construir_libro(ORIGEN, salida, dir_salida=tmp_path, catalogo=_catalogo_sintetico([n1, n2]))
    wb = load_workbook(salida)
    hojas = wb.sheetnames[5:]
    assert len(hojas) == 2 and len(set(hojas)) == 2 and all(len(h) <= 31 for h in hojas)
    enlaces = {r[1].value: r[1].hyperlink.location for r in wb["Tablas"].iter_rows(min_row=2) if r[1].value}
    assert enlaces[n1] == f"'{hojas[0]}'!A1" and enlaces[n2] == f"'{hojas[1]}'!A1"


def test_sin_vinculos_externos(libro_real):
    assert libro_real._external_links == []
    assert len(libro_real.defined_names) == 0


def test_fecha_es_fecha(libro_real):
    ws = libro_real[libro_real.sheetnames[5]]
    cab = [c.value for c in ws[1]]
    celda = ws.cell(2, cab.index("fecha") + 1)
    assert isinstance(celda.value, datetime) and celda.number_format == "yyyy-mm-dd"
    assert ws.cell(2, cab.index("cod_ccaa") + 1).number_format == "@"
    assert ws.cell(2, cab.index("cod_ccaa") + 1).value == "00"
    assert isinstance(ws.cell(2, cab.index("num_viviendas") + 1).value, (int, float))


def test_general_sin_kpis_ni_errores(libro_real):
    ws = libro_real["General"]
    cab = [c.value for c in ws[1]]
    k = cab.index("KPIs") + 1
    assert ws.cell(1, k).comment is not None and "vaciada" in ws.cell(1, k).comment.text
    assert all(ws.cell(i, k).value is None for i in range(2, ws.max_row + 1))
    errores = {"#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!"}
    assert not any(c.value in errores for fila in ws.iter_rows() for c in fila)


def test_tabla_grande_hoja_muestra(tmp_path, monkeypatch):
    monkeypatch.setattr(cl, "LIMITE_FILAS_HOJA", 50)
    monkeypatch.setattr(cl, "FILAS_MUESTRA", 10)
    _csv(tmp_path, "grande", 60)
    salida = tmp_path / "v2.xlsx"
    construir_libro(ORIGEN, salida, dir_salida=tmp_path, catalogo=_catalogo_sintetico(["grande"]))
    wb = load_workbook(salida)
    ws = wb["grande"]
    assert "etl/salida/grande.csv" in ws["A1"].value
    assert ws.max_row == 12            # nota + cabecera + 10 filas
    t = wb["Tablas"]
    cab = [c.value for c in t[1]]
    assert t.cell(2, cab.index("Filas") + 1).value == 60


def test_general_cuenta_errores_fuera_de_kpis(tmp_path, capsys):
    from openpyxl import Workbook
    src = Workbook()
    src.active.title = "General"
    src["General"].append(["Dataset ID", "Nota", "KPIs"])
    src["General"].append(["A", "#REF!", "#VALUE!"])
    src["General"].append(["B", "ok", "#VALUE!"])
    src.create_sheet("Maestro Territorios").append(["alias", "nivel", "cod_territorio", "cod_ccaa", "nombre_ccaa",
                                                    "cod_provincia", "nombre_provincia", "cod_municipio",
                                                    "nombre_municipio", "nombre_territorio", "fuente_alias",
                                                    "cod_fuente_alias"])
    _csv(tmp_path, "t", 2)
    ruta = tmp_path / "origen.xlsx"
    src.save(ruta)
    salida = tmp_path / "v2.xlsx"
    construir_libro(ruta, salida, dir_salida=tmp_path, catalogo=_catalogo_sintetico(["t"]))
    ws = load_workbook(salida)["General"]
    assert ws["B2"].value is None and ws["B3"].value == "ok"
    assert ws["C2"].value is None and ws["C3"].value is None
    out = capsys.readouterr().out
    assert "2 celdas con error en la columna KPIs" in out and "1 celdas con error de Excel vaciadas" in out
