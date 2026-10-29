"""MIVAU-SERPAVI: Sistema Estatal de Referencia del Precio del Alquiler de Vivienda, base de datos 2011-2024.

Tabla `mivau_serpavi_alquiler`, anual, niveles CCAA, Provincia, Municipio, Distrito y Seccion (censal), dimensión
`tipologia` (COLECTIVA / UNIFAMILIAR; el archivo no trae total de tipologías) y diez medidas: renta mensual por m2
(mediana, p25, p75), renta mensual de todo el inmueble (mediana, p25, p75), superficie (mediana, p25, p75) y
número de inmuebles del subconjunto de alquiler.

El archivo (~68 MB, .xlsx) tiene una hoja por nivel en formato ANCHO: una fila por territorio y 280 columnas = 14
años x 20 medidas, con nombres codificados (`BI_ALVHEPCO_TVC_11`, `ALQM2_LV_M_VC_11`…; `_AA` = año). Se lee con
openpyxl en modo `read_only` fila a fila; las columnas se interpretan por su nombre (`columnas_medidas`), y
cualquier nombre desconocido hace fallar el proceso.

Territorios: el código INE de la fila manda (CCAA, CPRO, CUMUN, CUDIS o CUSEC) y se busca por código en el
maestro (`territorio_por_codigo`), que incluye los distritos y secciones de etl/lib/secciones_serpavi.csv. Si las
columnas CPRO/CUMUN de la fila contradicen el código, solo se acepta en la lista cerrada `CODIGOS_INCOHERENTES`.
La fila sin código de la hoja CCAA (fila 21: solo recuentos 2012-2024, igual a la suma de las comunidades) no se
carga: no tiene etiqueta (pendiente de revisión por el usuario). Las filas (territorio, año, tipología) sin
ningún dato no entran.

Medianas y percentiles no se suman: `CUADRE = False`. Los recuentos sí: `comprobar_recuentos` comprueba en cada
ejecución CCAA = suma de provincias, provincia = suma de municipios, municipio = suma de distritos y distrito =
suma de secciones, con 35 diferencias conocidas de la fuente (`EXCEPCIONES_RECUENTOS`, lista cerrada).
"""
import re
import time
import warnings
from datetime import date
from pathlib import Path

import pandas as pd

from etl.lib import mivau
from etl.lib.fuente_comun import DESCARGAS, LIBRO, escribir_csv, url_dataset, validar_o_fallar
from etl.lib.modelo import construir_tabla
from etl.lib.secciones_serpavi_generar import normalizar_codigo
from etl.lib.territorios import cargar_maestro, territorio_por_codigo

DATASET_ID = "MIVAU-SERPAVI"
TABLA = "mivau_serpavi_alquiler"
DIMENSIONES = ["tipologia"]
RECUENTO = "num_inmuebles_alquiler"
MEDIDAS_DECIMALES = ["renta_mediana_eur_m2_mes", "renta_p25_eur_m2_mes", "renta_p75_eur_m2_mes",
                     "renta_mediana_eur_mes", "renta_p25_eur_mes", "renta_p75_eur_mes",
                     "superficie_mediana_m2", "superficie_p25_m2", "superficie_p75_m2"]
MEDIDAS = [*MEDIDAS_DECIMALES, RECUENTO]
CUADRE = False
TERRITORIO_EXTRA = ["cod_distrito", "cod_seccion"]
DECIMALES = 2

# Hoja -> (nivel, columnas de identificación esperadas, columna del código, longitud, prefijo de cod_territorio)
HOJAS = {
    "CCAA": ("CCAA", ["CCAA", "LITCCAA"], "CCAA", 2, "CA"),
    "Provincias": ("Provincia", ["CPRO", "LITPRO"], "CPRO", 2, "PR"),
    "Municipios": ("Municipio", ["CPRO", "NPRO", "CUMUN", "NMUN"], "CUMUN", 5, "MU"),
    "Distritos": ("Distrito", ["CPRO", "LITPRO", "CUMUN", "LITMUN", "CUDIS"], "CUDIS", 7, "DI"),
    "Secciones censales": ("Seccion", ["CPRO", "LITPRO", "CUMUN", "LITMUN", "CUSEC"], "CUSEC", 10, "SC"),
}
TIPOLOGIA = {"C": "COLECTIVA", "U": "UNIFAMILIAR"}
_ESTADISTICO = {"M": "mediana", "25": "p25", "75": "p75"}
# Familia de columnas -> plantilla del nombre de la medida ({e} = mediana / p25 / p75)
_RE_RECUENTO = re.compile(r"^BI_ALVHEPCO_TV([CU])_(\d{2})$")
_RE_MEDIDA = re.compile(r"^(ALQM2_LV|ALQTBID12|SLVM2)_(M|25|75)_V([CU])_(\d{2})$")
_FAMILIA = {"ALQM2_LV": "renta_{e}_eur_m2_mes", "ALQTBID12": "renta_{e}_eur_mes", "SLVM2": "superficie_{e}_m2"}

# Filas cuyo CPRO/CUMUN no coincide con su código (el código manda; se comprobó con los recuentos y el nombre).
# Código -> descripción. Lista cerrada: otra incoherencia hace fallar el proceso.
CODIGOS_INCOHERENTES = {
    "3100101": "Distritos fila 6146: CUMUN 31002 (Abárzuza) pero CUDIS de Abáigar (31001); recuentos = 31001",
    "0801908048": "Secciones fila 4706: CPRO '09', CUMUN 08020; el código es de Barcelona (08019)",
    "0801909055": "Secciones fila 4829: CPRO '09', CUMUN 08020; el código es de Barcelona (08019)",
    "1101208012": "Secciones fila 8576: CPRO '12', CUMUN 11013; el código es de Cádiz (11012)",
    "1402106037": "Secciones fila 10403: CUMUN 14022; el código es de Córdoba (14021)",
    "4109105045": "Secciones fila 29741: CPRO '42'; el código es de Sevilla (41091)",
}
# Recuentos no enteros de la fuente que pasan a NULL: (nivel, código INE, año, tipología).
RECUENTOS_NO_ENTEROS = {
    ("Municipio", "50275", 2022, "COLECTIVA"): "Valmadrid, Municipios HQ8072 = 103.383236…; distrito y sección = 0",
}
# Diferencias padre - suma de hijos en el número de inmuebles (descarga 2026-10-29), todas en Navarra:
# (nivel del padre, código INE del padre, año, tipología) -> diferencia. Pendiente de revisión por el usuario.
EXCEPCIONES_RECUENTOS = {
    ("Provincia", "31", 2023, "UNIFAMILIAR"): 2.0,
    ("Provincia", "31", 2024, "COLECTIVA"): 2.0,
    ("Provincia", "31", 2024, "UNIFAMILIAR"): 56.0,
    **{("Municipio", m, 2022, t): v for (m, t), v in {
        ("31134", "COLECTIVA"): -2.0, ("31137", "UNIFAMILIAR"): 1.0, ("31138", "COLECTIVA"): 12.0,
        ("31140", "COLECTIVA"): -44.0, ("31140", "UNIFAMILIAR"): -17.0, ("31142", "COLECTIVA"): 42.0,
        ("31142", "UNIFAMILIAR"): 15.0, ("31144", "UNIFAMILIAR"): 2.0, ("31147", "COLECTIVA"): -58.0,
        ("31149", "UNIFAMILIAR"): 1.0, ("31150", "UNIFAMILIAR"): -1.0, ("31152", "COLECTIVA"): -23.0,
        ("31152", "UNIFAMILIAR"): 1.0, ("31153", "COLECTIVA"): 29.0, ("31153", "UNIFAMILIAR"): 1.0}.items()},
    **{("Distrito", d, 2022, t): v for (d, t), v in {
        ("3113401", "COLECTIVA"): 2.0, ("3113701", "UNIFAMILIAR"): -1.0, ("3113801", "COLECTIVA"): -12.0,
        ("3114001", "COLECTIVA"): 44.0, ("3114001", "UNIFAMILIAR"): 17.0, ("3114201", "COLECTIVA"): -42.0,
        ("3114201", "UNIFAMILIAR"): -15.0, ("3114401", "UNIFAMILIAR"): -2.0, ("3114701", "COLECTIVA"): 58.0,
        ("3114901", "UNIFAMILIAR"): -1.0, ("3115001", "UNIFAMILIAR"): 1.0, ("3115201", "COLECTIVA"): 8.0,
        ("3115202", "COLECTIVA"): 15.0, ("3115202", "UNIFAMILIAR"): -1.0, ("3115301", "COLECTIVA"): -5.0,
        ("3115301", "UNIFAMILIAR"): 4.0, ("3115302", "UNIFAMILIAR"): -5.0}.items()},
}
# (nivel del padre, columna que agrupa a los hijos, nivel de los hijos)
_JERARQUIA = [("CCAA", "cod_ccaa", "Provincia"), ("Provincia", "cod_provincia", "Municipio"),
              ("Municipio", "cod_municipio", "Distrito"), ("Distrito", "cod_distrito", "Seccion")]


def columnas_medidas(cabecera, n_id: int) -> dict[int, tuple[int, str, str]]:
    """{índice de columna: (año, tipología, medida)} de las columnas tras las `n_id` de identificación.

    ValueError si un nombre no es de ninguna familia conocida o si a un (año, tipología) le falta alguna medida.
    """
    resultado = {}
    for j in range(n_id, len(cabecera)):
        nombre = str(cabecera[j]).strip() if cabecera[j] is not None else ""
        m = _RE_RECUENTO.match(nombre)
        if m:
            resultado[j] = (2000 + int(m.group(2)), TIPOLOGIA[m.group(1)], RECUENTO)
            continue
        m = _RE_MEDIDA.match(nombre)
        if not m:
            raise ValueError(f"{DATASET_ID}: columna {j + 1} con nombre desconocido {nombre!r}")
        medida = _FAMILIA[m.group(1)].format(e=_ESTADISTICO[m.group(2)])
        resultado[j] = (2000 + int(m.group(4)), TIPOLOGIA[m.group(3)], medida)
    bloques: dict[tuple, set] = {}
    for anio, tip, medida in resultado.values():
        bloques.setdefault((anio, tip), set()).add(medida)
    for clave, medidas in bloques.items():
        if medidas != set(MEDIDAS) or sum(1 for v in resultado.values() if v[:2] == clave) != len(MEDIDAS):
            raise ValueError(f"{DATASET_ID}: el bloque {clave} no tiene exactamente las {len(MEDIDAS)} medidas")
    return resultado


def comprobar_coherencia(hoja: str, codigo: str, ids: dict) -> bool:
    """True si CPRO/CUMUN de la fila coinciden con el código; False si no coinciden pero el código está en
    `CODIGOS_INCOHERENTES`; ValueError en otro caso."""
    esperado = {"CPRO": codigo[:2], "CUMUN": codigo[:5]}
    malos = {c: v for c, v in ids.items() if c in esperado and v != esperado[c]}
    if not malos:
        return True
    if codigo in CODIGOS_INCOHERENTES:
        return False
    raise ValueError(f"{DATASET_ID}: hoja {hoja!r}, código {codigo}: columnas {malos} no coinciden con el código "
                     f"(añadir a CODIGOS_INCOHERENTES solo tras revisarlo)")


def a_recuento(valor, clave: tuple):
    """Número de inmuebles como entero; None si vacío. Un no entero es error salvo en `RECUENTOS_NO_ENTEROS`
    (pasa a NULL)."""
    v = mivau.a_numero(valor)
    if v is None:
        return None
    if v != int(v):
        if clave in RECUENTOS_NO_ENTEROS:
            return None
        raise ValueError(f"{DATASET_ID}: recuento no entero {v!r} en {clave}")
    return int(v)


def leer_serpavi(path: Path) -> tuple[pd.DataFrame, list[str]]:
    """Archivo SERPAVI → filas largas (nivel, codigo, anio, tipologia, medidas) y avisos.

    Una fila por (territorio, año, tipología) con alguna medida; las celdas '' o vacías son NULL; los números
    guardados como texto ('600') se convierten.
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    columnas = {c: [] for c in ["hoja", "nivel", "codigo", "anio", "tipologia", *MEDIDAS]}
    avisos, incoherentes, no_enteros = [], [], []
    try:
        for hoja, (nivel, ids_esperados, col_codigo, largo, _) in HOJAS.items():
            filas = wb[hoja].iter_rows(values_only=True)
            cab = [str(c).strip() if c is not None else "" for c in next(filas)]
            n_id = len(ids_esperados)
            if cab[:n_id] != ids_esperados:
                raise ValueError(f"{DATASET_ID}: hoja {hoja!r}: columnas de identificación {cab[:n_id]}, "
                                 f"se esperaban {ids_esperados}")
            mapa = columnas_medidas(cab, n_id)
            bloques: dict[tuple, dict[str, int]] = {}
            for j, (anio, tip, medida) in mapa.items():
                bloques.setdefault((anio, tip), {})[medida] = j
            j_codigo = ids_esperados.index(col_codigo)
            sin_codigo = 0
            for i, f in enumerate(filas, 2):
                if f[j_codigo] is None:
                    if all(v in (None, "") for v in f):
                        continue
                    if hoja == "CCAA" and all(v is None for v in f[:n_id]) and sin_codigo == 0:
                        sin_codigo += 1  # fila 21: suma de las CCAA sin etiqueta (no se carga)
                        avisos.append(f"{DATASET_ID}: hoja 'CCAA' fila {i} sin código ni nombre (solo recuentos); "
                                      f"no se carga")
                        continue
                    raise ValueError(f"{DATASET_ID}: hoja {hoja!r} fila {i} con datos y sin código")
                codigo = normalizar_codigo(f[j_codigo], largo, f"{DATASET_ID}, hoja {hoja!r} fila {i}")
                if largo > 2:
                    ids = {"CPRO": normalizar_codigo(f[0], 2, f"{hoja} fila {i}")}
                    if nivel != "Municipio":
                        ids["CUMUN"] = normalizar_codigo(f[2], 5, f"{hoja} fila {i}")
                    if not comprobar_coherencia(hoja, codigo, ids):
                        incoherentes.append(codigo)
                for (anio, tip), cols in bloques.items():
                    valores = {m: f[j] for m, j in cols.items()}
                    if all(v in (None, "") for v in valores.values()):
                        continue
                    clave = (nivel, codigo, anio, tip)
                    recuento = a_recuento(valores[RECUENTO], clave)
                    if recuento is None and clave in RECUENTOS_NO_ENTEROS and valores[RECUENTO] not in (None, ""):
                        no_enteros.append(f"{clave} = {valores[RECUENTO]!r}")
                    decimales = {m: mivau.a_numero(valores[m]) for m in MEDIDAS_DECIMALES}
                    if recuento is None and all(v is None for v in decimales.values()):
                        continue
                    columnas["hoja"].append(hoja)
                    columnas["nivel"].append(nivel)
                    columnas["codigo"].append(codigo)
                    columnas["anio"].append(anio)
                    columnas["tipologia"].append(tip)
                    for m in MEDIDAS_DECIMALES:
                        v = decimales[m]
                        columnas[m].append(None if v is None else round(v, DECIMALES))
                    columnas[RECUENTO].append(recuento)
    finally:
        wb.close()
    if incoherentes:
        avisos.append(f"{DATASET_ID}: {len(incoherentes)} filas con CPRO/CUMUN distintos de su código (manda el "
                      f"código): {', '.join(sorted(set(incoherentes)))}")
    if no_enteros:
        avisos.append(f"{DATASET_ID}: recuentos no enteros pasados a NULL: {'; '.join(no_enteros)}")
    df = pd.DataFrame(columnas)
    df[RECUENTO] = df[RECUENTO].astype("Int64")
    for m in MEDIDAS_DECIMALES:
        df[m] = df[m].astype("float64")
    return df, avisos


def diferencias_recuentos(df: pd.DataFrame) -> dict[tuple, float]:
    """{(nivel del padre, código INE del padre, año, tipología): padre - suma de hijos} distintas de 0.

    Solo se comparan padres con recuento y con algún hijo con recuento (un hijo NULL cuenta como 0)."""
    vistas = {}
    v = df.assign(_n=pd.to_numeric(df[RECUENTO], errors="coerce").astype("float64"))
    for nivel_padre, clave, nivel_hijo in _JERARQUIA:
        padres = v[v["nivel"] == nivel_padre].set_index([clave, "anio", "tipologia"])["_n"].dropna()
        hijos = (v[v["nivel"] == nivel_hijo].groupby([clave, "anio", "tipologia"])["_n"].sum(min_count=1)
                 .dropna())
        dif = (padres - hijos).dropna()
        for (cod, anio, tip), d in dif[dif != 0].items():
            vistas[(nivel_padre, cod, int(anio), tip)] = float(d)
    return vistas


def comprobar_recuentos(df: pd.DataFrame) -> None:
    """ValueError si aparece una diferencia de recuentos fuera de `EXCEPCIONES_RECUENTOS` (o de otro importe);
    aviso si una conocida ya no aparece."""
    vistas = diferencias_recuentos(df)
    nuevas = {k: d for k, d in vistas.items() if EXCEPCIONES_RECUENTOS.get(k) != d}
    if nuevas:
        (nivel, cod, anio, tip), d = sorted(nuevas.items())[0]
        raise ValueError(f"{DATASET_ID}: el número de inmuebles de {nivel} {cod} {anio} {tip} no es la suma de sus "
                         f"hijos (diferencia {d:g}; {len(nuevas)} diferencias fuera de la lista cerrada)")
    desaparecidas = sorted(set(EXCEPCIONES_RECUENTOS) - set(vistas))
    if desaparecidas:
        warnings.warn(f"{DATASET_ID}: {len(desaparecidas)} diferencias de recuentos conocidas no aparecen (p. ej. "
                      f"{desaparecidas[0]})", stacklevel=2)


def transformar(path: Path, fecha_descarga: date, cuadre: bool = True) -> pd.DataFrame:
    """Lee el archivo, resuelve los territorios por código, comprueba los recuentos y valida la tabla.

    `cuadre=False` omite `comprobar_recuentos` (solo para archivos manipulados en tests)."""
    t0 = time.perf_counter()
    largo, avisos = leer_serpavi(path)
    for a in avisos:
        warnings.warn(a, stacklevel=2)
    maestro = cargar_maestro(LIBRO)
    prefijo = {nivel: p for nivel, _, _, _, p in HOJAS.values()}
    largo["cod_territorio"] = largo["nivel"].map(prefijo) + largo["codigo"]
    territorios = pd.DataFrame([territorio_por_codigo(maestro, c) for c in largo["cod_territorio"].unique()])
    if (territorios["nivel"].values != largo.drop_duplicates("cod_territorio")["nivel"].values).any():
        raise ValueError(f"{DATASET_ID}: un código del archivo corresponde a otro nivel en el maestro")
    df = largo.drop(columns=["nivel"]).merge(territorios, on="cod_territorio", how="left", validate="many_to_one")
    df["cod_distrito"] = df["codigo"].where(df["nivel"].isin(["Distrito", "Seccion"])).str[:7]
    df["cod_seccion"] = df["codigo"].where(df["nivel"] == "Seccion")
    if cuadre:
        comprobar_recuentos(df)
    df["periodo"] = df["anio"].astype(str)
    df["frecuencia"] = "A"
    df["trimestre"] = pd.array([pd.NA] * len(df), dtype="Int64")
    df["mes"] = pd.array([pd.NA] * len(df), dtype="Int64")
    fechas = {a: date(a, 1, 1) for a in df["anio"].unique()}
    df["fecha"] = df["anio"].map(fechas)
    df["provisional"] = 0
    tabla = construir_tabla(df, DIMENSIONES, MEDIDAS, fecha_descarga, territorio_extra=TERRITORIO_EXTRA)
    try:
        validar_o_fallar(tabla, TABLA, maestro)
    except ValueError as e:
        raise ValueError(f"{DATASET_ID}: {e}") from None
    transformar.segundos = round(time.perf_counter() - t0, 1)
    return tabla


def ejecutar(fecha_descarga: date) -> pd.DataFrame:
    """Descarga (~68 MB), transforma y escribe etl/salida/<TABLA>.csv (no versionado: supera los 50 MB)."""
    path = mivau.descargar(url_dataset(DATASET_ID), DATASET_ID, DESCARGAS)
    tabla = transformar(path, fecha_descarga)
    escribir_csv(tabla, TABLA)
    return tabla
