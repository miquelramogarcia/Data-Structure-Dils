"""Maestro de territorios: carga, resolución de nombres y expansión de uniprovinciales."""
import shutil
import tempfile
import unicodedata
from contextlib import contextmanager
from pathlib import Path

import pandas as pd

COLUMNAS_MAESTRO = ["alias", "nivel", "cod_territorio", "cod_ccaa", "nombre_ccaa", "cod_provincia",
                    "nombre_provincia", "cod_municipio", "nombre_municipio", "nombre_territorio",
                    "fuente_alias", "cod_fuente_alias"]
COLUMNAS_TERRITORIO = ["nivel", "cod_territorio", "cod_ccaa", "nombre_ccaa", "cod_provincia",
                       "nombre_provincia", "cod_municipio", "nombre_municipio"]
CCAA_UNIPROVINCIALES = {"03", "04", "06", "13", "14", "15", "17", "18", "19"}
EXTRA_CSV = Path(__file__).with_name("territorios_extra.csv")


class TerritorioDesconocido(ValueError):
    """El texto no corresponde a ningún territorio del maestro."""


def normalizar(txt: str) -> str:
    t = unicodedata.normalize("NFKD", str(txt)).encode("ascii", "ignore").decode()
    return " ".join(t.lower().split())


@contextmanager
def copia_temporal(xlsx: Path):
    """Entrega la ruta de una copia temporal del libro (el original no se toca ni se bloquea)."""
    with tempfile.TemporaryDirectory() as tmp:
        copia = Path(tmp) / "copia.xlsx"
        shutil.copy(xlsx, copia)
        yield copia


def leer_hoja_usuario(xlsx: Path, hoja: str, **kw) -> pd.DataFrame:
    """Lee una hoja del libro del usuario desde una copia temporal (el original no se toca)."""
    with copia_temporal(xlsx) as copia:
        return pd.read_excel(copia, sheet_name=hoja, **kw)


def leer_hoja_con_errores(xlsx: Path, hoja: str) -> pd.DataFrame:
    """Como leer_hoja_usuario, pero las celdas de error llegan como texto ('#VALUE!', '#REF!'…).

    pandas convierte esas celdas en NaN y no se pueden contar; openpyxl las entrega tal cual.
    Primera fila = cabecera; se descartan columnas sin cabecera y filas totalmente vacías.
    """
    from openpyxl import load_workbook
    with copia_temporal(xlsx) as copia:
        wb = load_workbook(copia, read_only=True, data_only=True)
        try:
            filas = [list(f) for f in wb[hoja].iter_rows(values_only=True)]
        finally:
            wb.close()
    cab = filas[0] if filas else []
    validas = [j for j, c in enumerate(cab) if c is not None]
    datos = [[f[j] if j < len(f) else None for j in validas] for f in filas[1:]]
    datos = [f for f in datos if any(v is not None for v in f)]
    return pd.DataFrame(datos, columns=[cab[j] for j in validas], dtype=object)


def cargar_maestro(xlsx: Path) -> pd.DataFrame:
    """Lee la hoja 'Maestro Territorios' desde una copia temporal y añade territorios_extra.csv."""
    df = leer_hoja_usuario(xlsx, "Maestro Territorios", header=None, dtype=str)
    if len(df) and str(df.iloc[0, 0]).strip().lower() == "alias":
        df = df.iloc[1:]
    df = df.iloc[:, :len(COLUMNAS_MAESTRO)].copy()
    df.columns = COLUMNAS_MAESTRO
    extra = pd.read_csv(EXTRA_CSV, dtype=str)
    if len(extra):
        df = pd.concat([df, extra[COLUMNAS_MAESTRO]], ignore_index=True)
    return df.reset_index(drop=True)


def resolver(maestro: pd.DataFrame, texto: str, nivel: str | None = None) -> dict:
    clave = normalizar(texto)
    m = maestro
    if nivel is not None:
        m = m[m["nivel"] == nivel]
    coincide = (m["alias"].map(lambda x: normalizar(x) if pd.notna(x) else None) == clave) | \
               (m["nombre_territorio"].map(lambda x: normalizar(x) if pd.notna(x) else None) == clave)
    enc = m[coincide]
    if enc.empty:
        raise TerritorioDesconocido(f"Territorio no encontrado en el maestro: {texto!r}")
    if len(enc) > 1 and nivel is None:
        ccaa = enc[enc["nivel"] == "CCAA"]
        if not ccaa.empty:
            enc = ccaa
    fila = enc.iloc[0]
    return {c: (None if pd.isna(fila[c]) else fila[c]) for c in COLUMNAS_TERRITORIO}


def _clave(fila: pd.Series, cols: list[str]) -> tuple:
    return tuple(None if pd.isna(fila[c]) else fila[c] for c in cols)


def expandir_uniprovinciales(df: pd.DataFrame, maestro: pd.DataFrame,
                             claves: list[str] | None = None) -> pd.DataFrame:
    """Añade la fila de provincia (copiando la de la CCAA) en las CCAA uniprovinciales si falta.

    La clave de deduplicación es la CCAA más `claves` (dimensiones + periodo). Si es None, se usan
    todas las columnas no territoriales que no sean numéricas (las medidas nunca entran en la clave).
    """
    if claves is None:
        otras = [c for c in df.columns if c not in COLUMNAS_TERRITORIO
                 and not pd.api.types.is_numeric_dtype(df[c])]
    else:
        otras = list(claves)
    existentes = {(f["cod_ccaa"],) + _clave(f, otras) for _, f in df[df["nivel"] == "Provincia"].iterrows()}
    nuevas = []
    for _, f in df[(df["nivel"] == "CCAA") & df["cod_ccaa"].isin(CCAA_UNIPROVINCIALES)].iterrows():
        k = (f["cod_ccaa"],) + _clave(f, otras)
        if k in existentes:
            continue
        prov = maestro[(maestro["nivel"] == "Provincia") & (maestro["cod_ccaa"] == f["cod_ccaa"])]
        if prov.empty:
            continue
        nueva = f.copy()
        for c in COLUMNAS_TERRITORIO:
            nueva[c] = prov.iloc[0][c]
        nuevas.append(nueva)
        existentes.add(k)
    if not nuevas:
        return df
    return pd.concat([df, pd.DataFrame(nuevas)], ignore_index=True)
