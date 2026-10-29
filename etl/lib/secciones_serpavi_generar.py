"""Genera etl/lib/secciones_serpavi.csv (distritos y secciones censales) a partir del archivo de SERPAVI.

Uso: python -m etl.lib.secciones_serpavi_generar [etl/descargas/MIVAU-SERPAVI_AAAAMMDD.xlsx]

Los códigos de distrito (CUDIS, 7 dígitos = CPRO + CMUN + DIST) y de sección (CUSEC, 10 dígitos = CPRO + CMUN +
DIST + SECC) son los del INE (seccionado del Censo de Población 2021, según la hoja Metadatos del archivo).
Se toman de las hojas "Distritos" y "Secciones censales" (y los distritos implícitos en los códigos de sección).
El código manda: el municipio sale de sus 5 primeros dígitos y se busca en etl/lib/municipios_ine.csv (INE a
1-1-2026), que da CCAA, provincia y nombres. Un código cuyo municipio no exista en el INE hace fallar el proceso,
salvo los cuatro municipios extinguidos de `MUNICIPIOS_EXTINGUIDOS` (solo aparecen, sin ningún dato, en las
hojas Municipios y Distritos): sus distritos no entran en el maestro.
"""
import re
import sys
from pathlib import Path

import pandas as pd

from etl.lib.territorios import COLUMNAS_MAESTRO, MUNICIPIOS_CSV, SECCIONES_CSV

RAIZ = Path(__file__).resolve().parents[2]
# Municipios fusionados antes de 2021 que SERPAVI conserva con filas vacías (código INE -> motivo).
MUNICIPIOS_EXTINGUIDOS = {
    "15026": "Cesuras: fusionado con Oza dos Ríos en Oza-Cesuras (15902) en 2013",
    "15063": "Oza dos Ríos: fusionado con Cesuras en Oza-Cesuras (15902) en 2013",
    "36011": "Cerdedo: fusionado con Cotobade en Cerdedo-Cotobade (36902) en 2016",
    "36012": "Cotobade: fusionado con Cerdedo en Cerdedo-Cotobade (36902) en 2016",
}
# Hoja -> (columna del código, longitud)
HOJAS_CODIGO = {"Distritos": ("CUDIS", 7), "Secciones censales": ("CUSEC", 10)}


def normalizar_codigo(valor, longitud: int, contexto: str) -> str:
    """Código INE como texto de `longitud` dígitos (algunas celdas vienen como número: 1001 -> '01001')."""
    if isinstance(valor, bool) or valor is None:
        raise ValueError(f"{contexto}: código vacío o no válido {valor!r}")
    if isinstance(valor, float) and valor == int(valor):
        valor = int(valor)
    txt = str(valor).strip()
    if isinstance(valor, int):
        txt = txt.zfill(longitud)
    if not re.fullmatch(rf"\d{{{longitud}}}", txt):
        raise ValueError(f"{contexto}: código {valor!r} no tiene {longitud} dígitos")
    return txt


def leer_codigos(path: Path) -> tuple[set[str], set[str]]:
    """(códigos de distrito, códigos de sección) de las hojas Distritos y Secciones censales."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        codigos = {}
        for hoja, (columna, largo) in HOJAS_CODIGO.items():
            filas = wb[hoja].iter_rows(values_only=True, max_col=5)
            cab = [str(c).strip() if c is not None else "" for c in next(filas)]
            if columna not in cab:
                raise ValueError(f"Hoja {hoja!r} sin columna {columna!r}: {cab}")
            j = cab.index(columna)
            codigos[hoja] = {normalizar_codigo(f[j], largo, f"{hoja}, fila {i}")
                             for i, f in enumerate(filas, 2) if any(v is not None for v in f)}
    finally:
        wb.close()
    secciones = codigos["Secciones censales"]
    distritos = codigos["Distritos"] | {s[:7] for s in secciones}
    return distritos, secciones


def generar(original: Path, destino: Path = SECCIONES_CSV) -> pd.DataFrame:
    """Escribe el CSV con las 12 columnas del maestro (nivel Distrito y Seccion); devuelve el DataFrame."""
    distritos, secciones = leer_codigos(original)
    mun = pd.read_csv(MUNICIPIOS_CSV, dtype=str).set_index("cod_municipio")
    filas = []
    for nivel, prefijo, codigos in (("Distrito", "DI", sorted(distritos)), ("Seccion", "SC", sorted(secciones))):
        for cod in codigos:
            cmun = cod[:5]
            if cmun not in mun.index:
                if cmun in MUNICIPIOS_EXTINGUIDOS and nivel == "Distrito":
                    continue
                raise ValueError(f"{nivel} {cod}: el municipio {cmun} no está en municipios_ine.csv")
            m = mun.loc[cmun]
            nombre = f"{m['nombre_municipio']} - distrito {cod[5:7]}"
            if nivel == "Seccion":
                nombre += f" - sección {cod[7:]}"
            filas.append({"alias": f"{prefijo}{cod}".lower(), "nivel": nivel, "cod_territorio": f"{prefijo}{cod}",
                          "cod_ccaa": m["cod_ccaa"], "nombre_ccaa": m["nombre_ccaa"],
                          "cod_provincia": m["cod_provincia"], "nombre_provincia": m["nombre_provincia"],
                          "cod_municipio": cmun, "nombre_municipio": m["nombre_municipio"],
                          "nombre_territorio": nombre, "fuente_alias": "MIVAU-SERPAVI",
                          "cod_fuente_alias": "CUDIS" if nivel == "Distrito" else "CUSEC"})
    df = pd.DataFrame(filas, columns=COLUMNAS_MAESTRO)
    df.to_csv(destino, index=False, encoding="utf-8")
    return df


if __name__ == "__main__":
    origen = Path(sys.argv[1]) if len(sys.argv) > 1 else max((RAIZ / "etl" / "descargas").glob("MIVAU-SERPAVI_*.xlsx"))
    resultado = generar(origen)
    print(f"{SECCIONES_CSV}: {(resultado['nivel'] == 'Distrito').sum()} distritos, "
          f"{(resultado['nivel'] == 'Seccion').sum()} secciones")
