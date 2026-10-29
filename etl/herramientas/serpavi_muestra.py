"""Extrae la muestra de SERPAVI que usan los tests (tests/fixtures/serpavi_muestra.xlsx) del archivo completo.

Uso: python -m etl.herramientas.serpavi_muestra etl/descargas/MIVAU-SERPAVI_AAAAMMDD.xlsx [destino]

El archivo original (~68 MB) no se versiona. La muestra conserva las seis hojas con sus nombres, la fila de
cabecera real y TODAS las columnas (los 14 años), pero solo las filas de `FILAS` (números de fila de Excel del
original, en el mismo orden). La hoja Metadatos se copia entera. La lista está documentada en
docs/fuentes/MIVAU-SERPAVI.md (qué territorio es cada fila y por qué está).
"""
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

RAIZ = Path(__file__).resolve().parents[2]
DESTINO = RAIZ / "tests" / "fixtures" / "serpavi_muestra.xlsx"


def _rango(a: int, b: int) -> list[int]:
    return list(range(a, b + 1))


# Hoja -> filas del original (sin la cabecera, que es siempre la fila 1 y se copia siempre).
FILAS = {
    "CCAA": _rango(2, 21),                      # las 19 comunidades y la fila 21 sin código (total nacional)
    "Provincias": [52, 53],                      # Ceuta (51) y Melilla (52)
    "Municipios": [2152,                         # Cesuras (15026): municipio extinguido, sin datos
                   4620, 4739,                   # Abáigar (31001) y Jaurrieta (31134, Navarra 2022)
                   8072,                         # Valmadrid (50275): recuento no entero en 2022
                   8099, 8100],                  # Ceuta y Melilla
    "Distritos": [6146,                          # 3100101 con CUMUN/LITMUN de 31002 (incoherente)
                  6290,                          # 3113401 (Navarra 2022: recuentos desplazados)
                  10460,                         # 5027501 (Valmadrid)
                  *_rango(10499, 10512)],        # los 6 distritos de Ceuta y los 8 de Melilla
    "Secciones censales": [8576,                 # 1101208012 con CPRO '12' (incoherente)
                           10403,                # 1402106037 con CPRO numérico y CUMUN 14022 (incoherente)
                           23450, 23643,         # 3100101001 y 3113401001
                           35673,                # 5027501001 (Valmadrid)
                           *_rango(36195, 36294)],  # las 56 secciones de Ceuta y las 44 de Melilla
}


def extraer(original: Path, destino: Path = DESTINO) -> None:
    wb = load_workbook(original, read_only=True, data_only=True)
    salida = Workbook()
    salida.remove(salida.active)
    try:
        meta = salida.create_sheet("Metadatos")
        for fila in wb["Metadatos"].iter_rows(values_only=True):
            meta.append(list(fila))
        for hoja, filas in FILAS.items():
            ws = salida.create_sheet(hoja)
            quiero = set(filas)
            tope = max(filas)
            copiadas = {}
            for i, fila in enumerate(wb[hoja].iter_rows(values_only=True, max_row=tope), 1):
                if i == 1:
                    ws.append(list(fila))
                elif i in quiero:
                    copiadas[i] = list(fila)
            faltan = quiero - set(copiadas)
            if faltan:
                raise ValueError(f"Hoja {hoja!r}: filas {sorted(faltan)} no existen en el original")
            for i in filas:
                ws.append(copiadas[i])
    finally:
        wb.close()
    salida.save(destino)


if __name__ == "__main__":
    extraer(Path(sys.argv[1]), Path(sys.argv[2]) if len(sys.argv) > 2 else DESTINO)
    print(f"Muestra escrita en {sys.argv[2] if len(sys.argv) > 2 else DESTINO}")
