"""PROTOTIPO: libro Excel de muestra de la carga manual de Idealista, con el mismo formato que Data Dils_v2.xlsx.

Reutiliza las funciones del constructor de producción (etl/construir_libro.py) para las hojas Tablas, Campos y
las hojas de datos (ocultas, muestra de hasta 1.000 filas). No toca Data Dils_v2.xlsx ni el catálogo.

Uso: python -m prototipos.carga_manual.libro   (antes: python -m prototipos.carga_manual.idealista)
"""
import pandas as pd
from openpyxl import Workbook

from etl.catalogo.sql import cargar_catalogo
from etl.construir_libro import _escribir, _hoja_datos, _hoja_tablas
from prototipos.carga_manual.idealista import AQUI, SALIDA, TIPOS

LIBRO_PROTOTIPO = AQUI / "idealista_prototipo.xlsx"


def construir_libro_prototipo(salida=LIBRO_PROTOTIPO):
    tablas, _ = cargar_catalogo()
    ids = [t["dataset_id"] for t in TIPOS.values()]
    tablas = tablas[tablas["dataset_id"].isin(ids)].copy()
    for t in TIPOS.values():          # las filas que tendría el catálogo si se adoptara
        i = tablas.index[tablas["dataset_id"] == t["dataset_id"]]
        tablas.loc[i, ["esquema", "tabla", "hoja", "formato_origen", "estado"]] = [
            "IDEALISTA", t["tabla"], t["tabla"], "Copia manual de la tabla HTML (prototipos/carga_manual/pegados/)",
            "con_datos"]
    campos = pd.concat([pd.read_csv(AQUI / f"campos_idealista_{tipo}.csv", dtype=str, keep_default_na=False)
                        for tipo in TIPOS], ignore_index=True)

    wb = Workbook()
    wb.remove(wb.active)
    hojas, filas = {}, {}
    for t in TIPOS.values():
        hojas[t["dataset_id"]] = t["tabla"]
        filas[t["dataset_id"]] = _hoja_datos(wb, t["tabla"], t["tabla"], campos[campos["tabla"] == t["tabla"]], SALIDA)
    _hoja_tablas(wb, tablas, campos, hojas, filas)
    _escribir(wb, "Campos", campos.rename(columns={
        "tabla": "TABLA", "campo": "Campo", "tipo_mysql": "Tipo MySQL", "nulo": "Nulo", "clave": "Clave",
        "descripcion": "Descripción", "ejemplo": "Ejemplo", "generado": "Generada como"}),
        superior=True, ajustar={"Descripción"})
    for t in TIPOS.values():          # en el prototipo, el CSV completo vive en prototipos/carga_manual/salida/
        ws = wb[t["tabla"]]
        ws["A1"] = str(ws["A1"].value).replace("etl/salida/", "prototipos/carga_manual/salida/")
    orden = ["Tablas", "Campos"] + [t["tabla"] for t in TIPOS.values()]
    wb._sheets = [wb[n] for n in orden]
    wb.active = 0
    wb.save(salida)
    return salida


if __name__ == "__main__":
    print(f"Libro escrito en {construir_libro_prototipo()}")
